import io
import re
import sys
from typing import Any

import pdfplumber
from docx import Document

from src.exception import CustomException

SECTION_HEADINGS = {
    "summary": ["summary", "professional summary", "profile", "objective", "career objective", "about me"],
    "skills": ["skills", "technical skills", "core competencies", "key skills", "technical expertise", "skills & technologies"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "career history", "internship", "internships", "work history"],
    "education": ["education", "academic background", "academic qualifications", "qualifications", "educational qualifications"],
    "projects": ["projects", "academic projects", "personal projects", "key projects", "project experience"],
    "certifications": ["certifications", "certificates", "licenses", "professional certifications"],
    "achievements": ["achievements", "awards", "honors", "accomplishments"],
}

JD_HEADINGS = {
    "summary": ["about the role", "job summary", "role summary", "overview", "summary"],
    "responsibilities": ["responsibilities", "what you'll do", "what you will do", "duties", "role responsibilities", "key responsibilities"],
    "requirements": ["requirements", "required qualifications", "basic qualifications", "must have", "required skills", "qualifications"],
    "preferred": ["preferred qualifications", "preferred skills", "nice to have", "good to have", "preferred", "bonus"],
    "education": ["education", "educational requirements", "academic requirements"],
}


def _normalize_heading(line: str) -> str:
    line = re.sub(r"[^a-zA-Z0-9&' ]", " ", line.lower())
    return re.sub(r"\s+", " ", line).strip()


class DocumentParser:
    def extract_text(self, file_bytes: bytes, filename: str) -> str:
        try:
            ext = filename.lower().rsplit(".", 1)[-1]
            if ext == "pdf":
                return self._extract_pdf(file_bytes)
            if ext == "docx":
                return self._extract_docx(file_bytes)
            if ext == "txt":
                return file_bytes.decode("utf-8", errors="ignore")
            raise ValueError(f"Unsupported file type: .{ext}. Use PDF, DOCX, or TXT.")
        except Exception as e:
            raise CustomException(e, sys)

    def inspect_document(self, file_bytes: bytes, filename: str) -> dict[str, Any]:
        """Return document-level signals used by the ATS analyzer."""
        ext = filename.lower().rsplit(".", 1)[-1]
        info: dict[str, Any] = {"file_type": ext, "file_size_kb": round(len(file_bytes) / 1024, 1)}
        try:
            if ext == "pdf":
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    info["page_count"] = len(pdf.pages)
                    extracted_pages = [p.extract_text() or "" for p in pdf.pages]
                    info["pages_with_text"] = sum(bool(x.strip()) for x in extracted_pages)
                    info["scanned_like"] = bool(pdf.pages) and info["pages_with_text"] == 0
                    info["table_count"] = sum(len(p.find_tables()) for p in pdf.pages)
                    info["image_count"] = sum(len(getattr(p, "images", []) or []) for p in pdf.pages)
                    # Layout heuristic: compare word x-coordinates when available.
                    info["possible_columns"] = any(self._page_has_columns(p) for p in pdf.pages)
                    info["column_detection_supported"] = True
            elif ext == "docx":
                doc = Document(io.BytesIO(file_bytes))
                info["page_count"] = None
                info["table_count"] = len(doc.tables)
                info["image_count"] = len(doc.inline_shapes)
                # DOCX column detection is not implemented.
                # Do not infer columns from page_width; that value is not a column signal.
                info["possible_columns"] = False
                info["column_detection_supported"] = False
            else:
                info.update({"page_count": None, "table_count": 0, "image_count": 0, "possible_columns": False, "column_detection_supported": False})
        except Exception as e:
            info["inspection_error"] = str(e)
        return info

    @staticmethod
    def _page_has_columns(page) -> bool:
        try:
            words = page.extract_words() or []
            if len(words) < 30:
                return False
            xs = sorted(float(w["x0"]) for w in words if "x0" in w)
            if not xs:
                return False
            gaps = [xs[i + 1] - xs[i] for i in range(len(xs) - 1)]
            return max(gaps) > max(80, (page.width or 600) * 0.12)
        except Exception:
            return False

    def _extract_pdf(self, file_bytes: bytes) -> str:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    def _extract_docx(self, file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    def split_sections(self, text: str) -> dict:
        try:
            lines = text.splitlines()
            sections = {key: [] for key in SECTION_HEADINGS}
            sections["other"] = []
            current = "other"
            detected = []
            for line in lines:
                normalized = _normalize_heading(line)
                matched = None
                if normalized and len(normalized) <= 60:
                    for section, headings in SECTION_HEADINGS.items():
                        if normalized in headings:
                            matched = section
                            break
                if matched:
                    current = matched
                    detected.append(matched)
                    continue
                sections[current].append(line)
            result = {k: "\n".join(v).strip() for k, v in sections.items()}
            result["full_text"] = text
            result["_headings_detected"] = bool(detected)
            result["_detected_section_names"] = sorted(set(detected))
            return result
        except Exception as e:
            raise CustomException(e, sys)

    def split_jd_sections(self, text: str) -> dict:
        sections = {key: [] for key in JD_HEADINGS}
        sections["other"] = []
        current = "other"
        for line in text.splitlines():
            normalized = _normalize_heading(line)
            matched = None
            if normalized and len(normalized) <= 70:
                for section, headings in JD_HEADINGS.items():
                    if normalized in headings:
                        matched = section
                        break
            if matched:
                current = matched
            else:
                sections[current].append(line)
        result = {k: "\n".join(v).strip() for k, v in sections.items()}
        result["full_text"] = text
        return result

    def extract_bullets(self, section_text: str) -> list[str]:
        bullets = []
        for line in section_text.splitlines():
            cleaned = re.sub(r"^[\-\u2022\u25aa\u2023\*\u2013\u2014>]+\s*", "", line.strip())
            cleaned = re.sub(r"^\d+[.)]\s*", "", cleaned)
            if len(cleaned) >= 20:
                bullets.append(cleaned)
        return bullets

    def extract_requirements(self, jd_text: str) -> list[str]:
        try:
            jd_sections = self.split_jd_sections(jd_text)
            ordered = []
            for key in ["requirements", "responsibilities", "preferred", "education", "summary", "other"]:
                ordered.extend(jd_sections.get(key, "").splitlines())
            requirements = []
            for line in ordered:
                cleaned = re.sub(r"^[\-\u2022\u25aa\u2023\*\u2013\u2014>]+\s*", "", line.strip())
                if len(cleaned) > 140:
                    requirements.extend(p.strip() for p in re.split(r"(?<=[.!?])\s+", cleaned) if len(p.strip()) >= 20 and len(p.split()) >= 4)
                elif len(cleaned) >= 10 and len(cleaned.split()) >= 4:
                    # require at least 4 words -- filters out bare skill-name bullets
                    # (e.g. "Scikit-learn", "Core Skills") which keyword-coverage
                    # already handles more reliably than sentence embeddings can
                    requirements.append(cleaned)
            seen, unique = set(), []
            for r in requirements:
                key = re.sub(r"\s+", " ", r.lower()).strip()
                if key not in seen:
                    seen.add(key)
                    unique.append(r)
            return unique[:40]
        except Exception as e:
            raise CustomException(e, sys)
