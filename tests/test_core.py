from src.components.requirement_classifier import RequirementClassifier
from src.components.skill_extractor import SkillExtractor
from src.components.suggestion_engine import SuggestionEngine
from src.components.document_parser import DocumentParser


def test_alias_normalization():
    skills = SkillExtractor().extract("Built APIs with Fast API and sklearn; deployed with PowerBI.")
    assert "fastapi" in skills
    assert "scikit learn" in skills
    assert "power bi" in skills


def test_required_language_beats_frequency():
    classifier = RequirementClassifier()
    assert classifier.classify_text("Experience with Docker is required.") == "must_have"
    assert classifier.classify_text("Kubernetes is a plus.") == "nice_to_have"


def test_gap_classification_uses_context():
    extractor = SkillExtractor()
    resume = extractor.extract("Python and SQL")
    jd = extractor.extract_with_frequency("Experience with AWS is required. Docker is a plus.")
    gap = SuggestionEngine().analyze_gap(resume, jd, "Experience with AWS is required. Docker is a plus.")
    assert any(x["skill"] == "aws" for x in gap["must_have_missing"])
    assert any(x["skill"] == "docker" for x in gap["nice_to_have_missing"])


def test_jd_requirement_extraction():
    parser = DocumentParser()
    reqs = parser.extract_requirements("Requirements:\n- Python experience is required.\n- Docker is a plus.")
    assert len(reqs) == 2


def test_optimization_plan_separates_skill_and_evidence_gaps():
    extractor = SkillExtractor()
    jd_text = "Requirements:\n- Python is required.\n- AWS is required.\n- Docker is a plus.\nResponsibilities:\n- Deploy production APIs to cloud infrastructure."
    resume = extractor.extract("Python and SQL. Built data analysis projects.")
    jd = extractor.extract_with_frequency(jd_text)
    gap = SuggestionEngine().analyze_gap(resume, jd, jd_text)
    coverage = [{
        "requirement": "Deploy production APIs to cloud infrastructure.",
        "best_score": 25.0,
        "best_bullet": "Built data analysis projects."
    }]
    plan = SuggestionEngine().build_optimization_plan(gap, coverage)
    assert any(x["skill"] == "aws" for x in plan["skill_gaps_to_close"])
    assert any(x["skill"] == "docker" for x in plan["skill_gaps_to_close"])
    assert plan["target_score_pct"] == 100.0
    assert plan["target_is_guaranteed"] is False


def test_docx_column_detection_is_explicitly_unsupported():
    from docx import Document
    import io
    parser = DocumentParser()
    doc = Document()
    doc.add_paragraph("John Doe")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python SQL")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Built APIs")
    buf = io.BytesIO()
    doc.save(buf)
    info = parser.inspect_document(buf.getvalue(), "resume.docx")
    assert info["possible_columns"] is False
    assert info["column_detection_supported"] is False
